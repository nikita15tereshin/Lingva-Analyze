import io

from src.parser.reader import extract_and_normalize, normalize_text


def test_extract_txt_from_path(tmp_path):
    f = tmp_path / "doc.txt"
    f.write_text("Die   Bundeswehr   übt.\nNeuer Absatz.", encoding="utf-8")
    text = extract_and_normalize(str(f), "doc.txt")
    assert text == "Die Bundeswehr übt. Neuer Absatz."


def test_extract_txt_from_stream():
    stream = io.BytesIO("Das ist ein Test.".encode("utf-8"))
    text = extract_and_normalize(stream, "doc.txt")
    assert text == "Das ist ein Test."


def test_extract_docx():
    import docx as python_docx

    d = python_docx.Document()
    d.add_paragraph("Erster Absatz.")
    d.add_paragraph("Zweiter Absatz mit Wort Bundesheer.")
    buf = io.BytesIO()
    d.save(buf)
    buf.seek(0)

    text = extract_and_normalize(buf, "doc.docx")
    assert "Bundesheer" in text


def test_unsupported_format():
    try:
        extract_and_normalize(io.BytesIO(b""), "file.odt")
        assert False, "ожидалось ValueError"
    except ValueError:
        pass


def test_normalize_keeps_capitalization():
    assert normalize_text("Das  Gymnasium") == "Das Gymnasium"
